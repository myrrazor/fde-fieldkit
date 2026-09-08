from __future__ import annotations

import json
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, IO
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook

_EXTENSIONS = {
    ".csv": "csv",
    ".tsv": "tsv",
    ".xlsx": "xlsx",
    ".json": "json",
    ".jsonl": "jsonl",
    ".ndjson": "jsonl",
}
_FORMATS = set(_EXTENSIONS.values())
_NA_LITERALS = {"", "N/A", "null", "NULL", "None", "nan"}
MAX_DATASET_BYTES = 50 * 1024 * 1024
_NUMERIC_LITERAL = re.compile(r"-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")


@dataclass
class LoadedTable:
    """A string-preserving table plus details about how it was loaded."""

    df: pd.DataFrame
    fmt: str
    source: str
    warnings: list[str]


def detect_format(filename: str | None, sample: bytes) -> str:
    """Infer a supported table format from a filename or a small byte sample."""

    if filename:
        suffix = Path(filename).suffix.lower()
        if suffix in _EXTENSIONS:
            return _EXTENSIONS[suffix]

    head = sample[:4096]
    if head.startswith(b"PK"):
        return "xlsx"

    text = head.decode("utf-8-sig", errors="ignore").strip()
    if text.startswith("["):
        return "json"
    if text.startswith("{"):
        lines = [line for line in text.splitlines() if line.strip()]
        if len(lines) > 1:
            try:
                if all(isinstance(json.loads(line), dict) for line in lines):
                    return "jsonl"
            except json.JSONDecodeError:
                pass
        return "json"

    tabs = text.count("\t")
    commas = text.count(",")
    if tabs or commas:
        return "tsv" if tabs > commas else "csv"

    label = filename or "<buffer>"
    raise ValueError(f"can't detect a format for {label!r} — pass fmt=")


def load_table(
    src: Path | BinaryIO,
    *,
    filename: str | None = None,
    fmt: str | None = None,
    sheet: str | None = None,
) -> LoadedTable:
    """Load a path or binary stream without inferring application-level types."""

    raw, source = _read_source(src, filename)
    table_fmt = _clean_format(fmt) if fmt else detect_format(filename or source, raw)
    warnings: list[str] = []

    if table_fmt in {"csv", "tsv"}:
        df = pd.read_csv(
            BytesIO(raw),
            sep="\t" if table_fmt == "tsv" else ",",
            dtype=str,
            keep_default_na=False,
            na_filter=False,
        )
    elif table_fmt == "xlsx":
        df, warnings = _read_xlsx(raw, sheet)
    elif table_fmt == "json":
        df = _read_json(raw)
    elif table_fmt == "jsonl":
        df = _read_jsonl(raw)
    else:  # pragma: no cover - _clean_format owns this guard
        raise ValueError(f"unsupported table format: {table_fmt}")

    return LoadedTable(_as_nullable_strings(df), table_fmt, source, warnings)


def write_table(df: pd.DataFrame, dest: Path | IO[str] | IO[bytes], fmt: str) -> None:
    """Write a table, neutralizing spreadsheet formulas in CSV, TSV, and XLSX."""

    table_fmt = _clean_format(fmt)
    if table_fmt in {"csv", "tsv", "xlsx"}:
        df = _formula_safe_frame(df)
    if table_fmt == "xlsx":
        df.to_excel(dest, index=False, engine="openpyxl")
        return

    if table_fmt in {"csv", "tsv"}:
        text = df.to_csv(index=False, sep="\t" if table_fmt == "tsv" else ",")
    elif table_fmt == "json":
        text = df.to_json(orient="records", indent=2, date_format="iso")
    elif table_fmt == "jsonl":
        text = df.to_json(orient="records", lines=True, date_format="iso")
    else:  # pragma: no cover - _clean_format owns this guard
        raise ValueError(f"unsupported table format: {table_fmt}")
    _write_text(dest, text)


def _read_source(src: Path | BinaryIO, filename: str | None) -> tuple[bytes, str]:
    if isinstance(src, Path):
        if src.stat().st_size > MAX_DATASET_BYTES:
            raise ValueError("dataset exceeds the 50 MB total size limit")
        return src.read_bytes(), src.name

    raw = src.read(MAX_DATASET_BYTES + 1)
    if not isinstance(raw, bytes):
        raise TypeError("load_table expects a binary file object")
    if len(raw) > MAX_DATASET_BYTES:
        raise ValueError("dataset exceeds the 50 MB total size limit")
    return raw, filename or "<buffer>"


def _clean_format(fmt: str) -> str:
    cleaned = fmt.lower().removeprefix(".")
    if cleaned not in _FORMATS:
        choices = ", ".join(sorted(_FORMATS))
        raise ValueError(f"unsupported table format {fmt!r}; choose one of: {choices}")
    return cleaned


def _read_xlsx(raw: bytes, sheet: str | None) -> tuple[pd.DataFrame, list[str]]:
    _validate_xlsx_archive(raw)
    workbook = load_workbook(BytesIO(raw), data_only=True, read_only=False)
    if sheet is not None and sheet not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValueError(f"sheet {sheet!r} not found; available sheets: {available}")
    worksheet = workbook[sheet] if sheet else workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame(), []

    headers = [str(value) if value is not None else "" for value in rows[0]]
    warnings = []
    if worksheet.merged_cells.ranges:
        warnings.append(f"sheet {worksheet.title!r} contains merged cells")
    return pd.DataFrame(rows[1:], columns=headers), warnings


def _read_json(raw: bytes) -> pd.DataFrame:
    try:
        payload = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"invalid JSON: {exc}") from exc
    if not isinstance(payload, list):
        raise ValueError("JSON tables need a top-level list of records, not an object")
    if any(not isinstance(row, dict) for row in payload):
        raise ValueError("JSON tables need a top-level list of record objects")
    return pd.json_normalize(payload, max_level=1)


def _read_jsonl(raw: bytes) -> pd.DataFrame:
    records = []
    try:
        for line in raw.decode("utf-8-sig").splitlines():
            if not line.strip():
                continue
            record = json.loads(line)
            if not isinstance(record, dict):
                raise ValueError("JSONL rows must be objects")
            records.append(record)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"invalid JSONL: {exc}") from exc
    return pd.json_normalize(records, max_level=1)


def _as_nullable_strings(df: pd.DataFrame) -> pd.DataFrame:
    def clean(value: object) -> object:
        if value is None or value is pd.NA:
            return pd.NA
        text = str(value)
        return pd.NA if text in _NA_LITERALS else text

    return df.map(clean).astype("string")


def _formula_safe_frame(df: pd.DataFrame) -> pd.DataFrame:
    safe = df.map(formula_safe_value)
    safe.columns = [formula_safe_value(str(column)) for column in df.columns]
    return safe


def formula_safe_value(value: object) -> object:
    """Neutralize strings that spreadsheet programs could execute as formulas."""

    if not isinstance(value, str):
        return value
    candidate = value.lstrip(" ")
    if not candidate:
        return value
    first = candidate[0]
    dangerous = first in {"=", "+", "@", "\t", "\r"}
    if first == "-" and not _NUMERIC_LITERAL.fullmatch(candidate):
        dangerous = True
    return f"'{value}" if dangerous else value


def _validate_xlsx_archive(raw: bytes) -> None:
    try:
        with ZipFile(BytesIO(raw)) as archive:
            expanded_bytes = sum(item.file_size for item in archive.infolist())
    except BadZipFile as exc:
        raise ValueError("invalid XLSX archive") from exc
    if expanded_bytes > MAX_DATASET_BYTES:
        raise ValueError("XLSX expands beyond the 50 MB total dataset limit")


def _write_text(dest: Path | IO[str] | IO[bytes], text: str) -> None:
    if isinstance(dest, Path):
        dest.write_text(text, encoding="utf-8")
        return
    try:
        dest.write(text)
    except TypeError:
        dest.write(text.encode("utf-8"))
